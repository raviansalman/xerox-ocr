"""Throughput and latency benchmark against a running Document Intelligence API.

Generates a synthetic corpus (native PDFs, Word files, text files and a share of scanned PDFs), uploads it through
the HTTP API, waits until every document is processed, then measures query latency for each kind of question.

    python scripts/load_test.py --url http://localhost:8000 --key $DOCINTEL_KEY --count 2000 --scanned 0.1

All names and numbers in the generated documents are fictitious. Upload a corpus to a dedicated tenant: the
script does not delete what it uploads.
"""
from __future__ import annotations

import argparse
import io
import json
import random
import statistics
import sys
import tempfile
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import httpx

VENDORS = ["Northwind Traders", "Contoso Ltd", "Fabrikam Inc", "Tailspin Toys", "Wide World Importers", "Adventure Works",
           "Litware Inc", "Proseware Inc", "Woodgrove Bank", "Blue Yonder Airlines", "Coho Winery", "Lucerne Publishing"]
CITIES = ["Riyadh", "Jeddah", "Dubai", "London", "Chicago", "Toronto", "Berlin", "Singapore"]
STATES = ["California", "Texas", "New York", "Delaware", "Florida"]
SERVICES = ["managed print services", "toner supply", "printer maintenance", "document scanning", "fleet monitoring"]


def _invoice(i: int, r: random.Random) -> tuple[str, list[str]]:
    v, amt = r.choice(VENDORS), r.randint(1_000, 900_000)
    ident = f"INV-{2024 + (i // 3) % 3}-{i:05d}"
    return f"Invoice_{ident}", [
        "TAX INVOICE", f"Invoice No: {ident}", f"Invoice Date: {1 + i % 28} March {2024 + i % 3}",
        f"Bill to: {v}", f"Description: {r.choice(SERVICES)} for the {r.choice(CITIES)} office",
        f"Total Amount Due: USD {amt:,}.00", "Payment terms: 30 days"]


def _contract(i: int, r: random.Random) -> tuple[str, list[str]]:
    v, st = r.choice(VENDORS), r.choice(STATES)
    return f"Service_Agreement_{i:05d}", [
        "SERVICE AGREEMENT", f"Contract No. SA-{i:05d}", f"This agreement is made between Acme Corp and {v}.",
        f"The supplier provides {r.choice(SERVICES)} at all {r.choice(CITIES)} sites.",
        f"The agreement expires on 31 December {2026 + i % 4}.",
        f"This agreement is governed by the laws of the State of {st}.",
        "Either party may terminate this agreement with 60 days written notice.",
        f"Contract value: USD {r.randint(10, 900) * 1000:,}.00", "Signed by: Jane Doe"]


def _memo(i: int, r: random.Random) -> tuple[str, list[str]]:
    return f"Memo_{i:05d}", [
        "INTERNAL MEMO", f"Subject: {r.choice(SERVICES)} update {i}",
        f"The {r.choice(CITIES)} team completed the quarterly review of {r.choice(SERVICES)}.",
        f"Ticket reference TK-{i:06d} was closed after the printer fleet audit."]


KINDS = [_invoice, _contract, _memo]


def generate(folder: Path, count: int, scanned: float, seed: int = 7) -> list[Path]:
    import docx
    import fitz
    from PIL import Image, ImageDraw, ImageFont

    r = random.Random(seed)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 34)
    except OSError:
        font = ImageFont.load_default()
    paths = []
    for i in range(count):
        name, lines = KINDS[i % len(KINDS)](i, r)
        roll = r.random()
        if roll < scanned:
            img = Image.new("L", (1654, 1100), 255)
            d = ImageDraw.Draw(img)
            for n, line in enumerate(lines):
                d.text((120, 90 + n * 70), line, font=font, fill=0)
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            doc = fitz.open()
            doc.new_page(width=595, height=842).insert_image(fitz.Rect(0, 0, 595, 396), stream=buf.getvalue())
            p = folder / f"{name}_scan.pdf"
            doc.save(str(p))
        elif roll < scanned + (1 - scanned) / 3:
            d = docx.Document()
            d.add_heading(lines[0], level=1)
            for line in lines[1:]:
                d.add_paragraph(line)
            p = folder / f"{name}.docx"
            d.save(str(p))
        elif roll < scanned + 2 * (1 - scanned) / 3:
            p = folder / f"{name}.txt"
            p.write_text("\n".join(lines) + "\n", encoding="utf-8")
        else:
            doc = fitz.open()
            page = doc.new_page(width=595, height=842)
            for n, line in enumerate(lines):
                page.insert_text((60, 72 + n * 18), line, fontsize=11)
            p = folder / f"{name}.pdf"
            doc.save(str(p))
        paths.append(p)
    return paths


# Every identifier below exists in a generated corpus of 500 or more documents
# (invoices are documents 0, 3, 6 ..., agreements 1, 4, 7 ..., memos 2, 5, 8 ...).
QUERIES = {
    "identifier": ["INV-2026-00033", "SA-00118", "TK-000458"],
    "exact_phrase": ['"either party may terminate"', '"quarterly review"'],
    "keyword": ["toner supply Jeddah", "printer maintenance Dubai"],
    "typo": ["Northwnd Traders", "Fabrikm invoice"],
    "semantic": ["who can end the agreement early", "bills for print services"],
    "contextual": ["Can either party cancel the agreement?"],
    "count": ["How many invoices do we have?", "How many contracts are governed by Texas law?"],
    "aggregate": ["What is the total value of all invoices?", "How many documents per type?"],
    "filtered": ["contracts expiring in 2027", "invoices above USD 500,000"],
    "lookup": ["When does service agreement SA-00118 expire?"],
}


def pct(values: list[float], p: float) -> float:
    s = sorted(values)
    return round(s[min(len(s) - 1, round(p / 100 * (len(s) - 1)))], 1)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--url", default="http://localhost:8000")
    ap.add_argument("--key", required=True, help="API key with the uploader role for the benchmark tenant")
    ap.add_argument("--count", type=int, default=1000)
    ap.add_argument("--scanned", type=float, default=0.1, help="share of scanned (OCR) documents")
    ap.add_argument("--batch", type=int, default=25)
    ap.add_argument("--upload-concurrency", type=int, default=4)
    ap.add_argument("--query-rounds", type=int, default=20)
    ap.add_argument("--query-concurrency", type=int, default=8)
    ap.add_argument("--timeout", type=int, default=7200, help="seconds to wait for processing")
    ap.add_argument("--queries-only", action="store_true", help="skip generation and upload; measure queries only")
    a = ap.parse_args()

    h = {"X-API-Key": a.key}
    client = httpx.Client(base_url=a.url, headers=h, timeout=300)
    before = client.get("/api/v1/stats").raise_for_status().json()
    report: dict = {"count": a.count, "scanned_share": a.scanned}

    if a.queries_only:
        report.update(count=0, scanned_share=None)
        return _queries(client, a, report)
    with tempfile.TemporaryDirectory() as tmp:
        t = time.perf_counter()
        paths = generate(Path(tmp), a.count, a.scanned)
        report["generate_s"] = round(time.perf_counter() - t, 1)

        def upload(chunk: list[Path]) -> int:
            files = [("files", (p.name, p.read_bytes())) for p in chunk]
            out = httpx.post(f"{a.url}/api/v1/documents", headers=h, files=files, timeout=600).raise_for_status().json()
            return sum(1 for d in out["documents"] if d.get("status") != "rejected" and not d.get("duplicate"))

        t0 = time.perf_counter()
        batches = [paths[i:i + a.batch] for i in range(0, len(paths), a.batch)]
        with ThreadPoolExecutor(a.upload_concurrency) as pool:
            accepted = sum(pool.map(upload, batches))
        report["upload_s"] = round(time.perf_counter() - t0, 1)
        report["accepted"] = accepted

    target = before.get("indexed_documents", 0) + accepted
    failed0 = (before.get("by_status") or {}).get("failed", 0)
    while True:
        s = client.get("/api/v1/stats").raise_for_status().json()
        done = s["indexed_documents"] + (s.get("by_status") or {}).get("failed", 0) - failed0
        if done >= target:
            break
        if time.perf_counter() - t0 > a.timeout:
            print(json.dumps({"error": "timeout waiting for processing", "stats": s}), file=sys.stderr)
            return 1
        time.sleep(2)
    total_s = time.perf_counter() - t0
    report.update({
        "processing_s": round(total_s, 1),
        "docs_per_minute": round(accepted / total_s * 60, 1),
        "failed": (s.get("by_status") or {}).get("failed", 0) - failed0,
        "indexed_documents_total": s["indexed_documents"],
        "chunks_total": s.get("chunks"),
    })

    return _queries(client, a, report)


def _queries(client: httpx.Client, a, report: dict) -> int:
    def run(q: str) -> tuple[float, int]:
        t = time.perf_counter()
        r = client.post("/api/v1/query", json={"q": q, "limit": 10})
        r.raise_for_status()
        return (time.perf_counter() - t) * 1000, r.json()["total"]

    latency: dict = {}
    with ThreadPoolExecutor(a.query_concurrency) as pool:
        for kind, qs in QUERIES.items():
            for q in qs:                                  # warm-up
                run(q)
            res = list(pool.map(run, [q for _ in range(a.query_rounds) for q in qs]))
            ms = [x for x, _ in res]
            latency[kind] = {"p50_ms": pct(ms, 50), "p95_ms": pct(ms, 95), "mean_ms": round(statistics.mean(ms), 1),
                             "results": res[0][1]}
    report["query_latency"] = latency
    report["query_concurrency"] = a.query_concurrency
    print(json.dumps(report, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
