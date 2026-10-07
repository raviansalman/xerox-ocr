"""Deterministic test documents built with libraries the app already depends on."""
import io
from pathlib import Path
from typing import List, Sequence

import fitz
from PIL import Image, ImageDraw, ImageFont

FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"

PRESS_RELEASE_PAGES = [
    [
        "FOR IMMEDIATE RELEASE",
        "StorageChain Announces Decentralized Archive Service",
        "Austin, Texas - June 3, 2024 - StorageChain today announced a new archive",
        "service for enterprise records. The press release outlines pricing and availability.",
        "Media contact: Lisa Riordan, Director of Communications",
    ],
    [
        "About StorageChain",
        "StorageChain builds secure storage for regulated industries.",
        "For more information contact Lisa Riordan at the press office.",
    ],
]


def text_pdf(path: Path, pages: Sequence[Sequence[str]]) -> Path:
    doc = fitz.open()
    for lines in pages:
        page = doc.new_page(width=595, height=842)
        y = 72
        for line in lines:
            page.insert_text((60, y), line, fontsize=11)
            y += 16
    doc.save(str(path))
    return path


def numbered_report_pdf(path: Path, n_pages: int) -> Path:
    return text_pdf(path, [
        [f"Annual maintenance report page {i}", f"Section {i} covers printer fleet uptime and toner usage."]
        for i in range(1, n_pages + 1)
    ])


def render_page(lines: List[str], size_pt: float = 12, rtl: bool = False, dpi: int = 300,
                width_in: float = 8.27, height_in: float = 3.0) -> Image.Image:
    """Render lines as a grayscale 'scan'. Height is kept small so tests stay fast."""
    w, h = int(width_in * dpi), int(height_in * dpi)
    img = Image.new("L", (w, h), 255)
    d = ImageDraw.Draw(img)
    px = int(size_pt * dpi / 72)
    font = ImageFont.truetype(FONT, px)
    y = int(0.3 * dpi)
    for line in lines:
        if rtl:
            d.text((w - int(0.6 * dpi), y), line, font=font, fill=0, anchor="ra", direction="rtl", language="ar")
        else:
            d.text((int(0.6 * dpi), y), line, font=font, fill=0)
        y += int(px * 1.6)
    return img


def mixed_digital_plus_scan_pdf(path: Path, scan_lines: List[str]) -> Path:
    """Three text pages (well above the 30 words/page heuristic) followed by one image-only page."""
    doc = fitz.open()
    for p in range(3):
        page = doc.new_page(width=595, height=842)
        y = 72
        for j in range(12):
            page.insert_text((60, y), f"Digital contract page {p + 1} clause {j + 1} governs service levels.", fontsize=10)
            y += 15
    buf = io.BytesIO()
    render_page(scan_lines, size_pt=14, height_in=11.69).save(buf, format="PNG")
    page = doc.new_page(width=595, height=842)
    page.insert_image(page.rect, stream=buf.getvalue())
    doc.save(str(path))
    return path


def invoice_docx(path: Path) -> Path:
    from docx import Document

    d = Document()
    d.add_paragraph("INVOICE 2024-0042")
    d.add_paragraph("Bill to: Gulf Trading LLC, Dubai")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Item"
    t.cell(0, 1).text = "Amount"
    t.cell(1, 0).text = "Managed print services"
    t.cell(1, 1).text = "USD 12,500"
    d.add_paragraph("Payment due within 30 days.")
    d.save(str(path))
    return path
